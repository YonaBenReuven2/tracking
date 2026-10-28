#!/usr/bin/env python3
"""Convert the old Tracking.xlsx spreadsheet into a log.json for the practice log app.

Usage:
    python tools/import_tracking.py "path/to/Tracking.xlsx" -o import/log.json

Then open the app, go to Settings and use "Load a log file" (or "Start over from a
file") with the result.

The spreadsheet is free text, so the script interprets every cell. It prints a report of
everything it interpreted, so you can check it before loading the result. The rules are
the tables right below; edit them and re-run if something comes out wrong.

How a row is read:
  * Column B..H hold what was done. The cells from B up to the first empty cell are your
    own practice; anything after an empty cell, or a row that only starts at E or later,
    was done at a lesson.
  * Breathing: "sfch" and "sfch no random" = sfch in set order, "all breathing" =
    set and random order, "sfch songs" = the rhythmic one, "5 sfch" = an older
    variation of sfch. A number of seconds ("35s") is the separate holding-breath
    exercise, its top hold (10, 15, 20 ... 35).
  * "exercise text" = that exercise with "text" as its variation.

Needs openpyxl (pip install openpyxl).
"""
import argparse
import datetime as dt
import json
import re
import sys
from collections import OrderedDict, defaultdict

# ---------------------------------------------------------------------------
# Rules. Edit these and re-run.
# ---------------------------------------------------------------------------

CATEGORIES = [("breathing", "Breathing"), ("exercises", "Exercises")]

SFCH = "sfch"
HOLD = "holding breath"
BREATHING = [SFCH, HOLD]

# Exercises folded into another one as a variation: old name -> (exercise, variation).
COMBINE = {
    "evil laugh 1": ("evil laugh", "1"),
    "evil laugh 2": ("evil laugh", "2"),
    "evil laugh 3": ("evil laugh", "3"),
    "hallelujah 3": ("kiki", "hallelujah 3"),
    "nasal hum": ("hum", "nasal"),
    "mimema": ("vavevivovu", "mimema"),
    "va": ("vavevivovu", "va"),
    "vavevivovu - a": ("vavevivovu", "a"),
    "dynamic vi": ("gliss", "octave"),
}

# Each exercise's own list of variations, in the order the app shows them.
VARIATIONS = {
    SFCH: ["set order", "random order", "songs"],
    "evil laugh": ["1", "2", "3"],
    "kiki": ["hallelujah 3"],
    "hum": ["nasal", "g"],
    "vavevivovu": ["a", "mimema", "va"],
    "gliss": ["compound", "octave", "from the top"],
}

# An entry with no variation that really means a particular one.
PLAIN_MEANS = {"gliss": "compound"}

# Exercises that were never in the right-hand pool but appear in the log.
EXTRA_EXERCISES = ["vivave"]

# Anything in the pool during the last RECENT_WEEKS weeks is on the board;
# everything else is archived.
RECENT_WEEKS = 4

# Typos and shorthand, replaced at the start of a line (whole words only).
PREFIX_ALIASES = {
    "all beathing": "all breathing",
    "breathing": "all breathing",
    "no random": "sfch no random",
    "f sfch": "5 sfch",
    "loius": "louis",
    "looee": "louis",
    "vibratto": "vibrato",
    "new gliss": "gliss new",
    "mouth hum": "hum mouth",
    "brr gliss": "gliss brr",
    "brr up down up": "up down up brr",
}

# Rows where the lesson part follows the own part with no empty cell between:
# date -> column where the lesson starts (B=0, C=1, D=2, E=3, F=4, G=5, H=6).
LESSON_STARTS = {
    "2025-07-09": 4,  # gliss, ma_a_a, gliss
    "2025-07-16": 4,  # gliss, do
}

# Whole cells that need a hand-made reading: exact cell text -> list of
# (exercise, [variations], note). An empty list ignores the cell.
CELL_OVERRIDES = {
    "yo": [("yo eee", [], "")],
    "none": [],
}

# Exercise ids for names that don't slugify into something readable.
ID_OVERRIDES = {"א-ת": "alef-tav"}

# ---------------------------------------------------------------------------

DAY_COL = 1               # A: date
LOG_COLS = range(2, 9)    # B..H: what was done that day
POOL_COLS = range(9, 27)  # I..Z: the rotation grid
SECONDS = re.compile(r"^(?:then\s+)?\d+\s*s\b", re.I)
KEY_ORDER = ["id", "date", "exercise", "name", "category", "variations", "lesson", "archived", "order",
             "label", "note", "deleted", "updated"]


class Report:
    def __init__(self):
        self.items = defaultdict(list)

    def add(self, kind, date, text, result):
        self.items[kind].append((date, text, result))

    def render(self):
        out = []
        for kind, rows in self.items.items():
            out.append(f"\n{kind} ({len(rows)})")
            seen = OrderedDict()
            for date, text, result in rows:
                key = (text, result)
                if key in seen:
                    seen[key][1] += 1
                else:
                    seen[key] = [date, 1]
            for (text, result), (date, n) in seen.items():
                when = f"{n}× from {date}" if n > 1 else date
                out.append(f"  {when}: {text!r} -> {result}")
        return "\n".join(out)


def clean(text):
    return re.sub(r"\s+", " ", str(text)).strip()


def slug(name):
    if name in ID_OVERRIDES:
        return ID_OVERRIDES[name]
    return re.sub(r"[^a-z0-9]+", "-", name.lower()).strip("-") or "ex"


def apply_aliases(line):
    for short, full in sorted(PREFIX_ALIASES.items(), key=lambda kv: -len(kv[0])):
        if line == short or line.startswith(short + " ") or line.startswith(short + ","):
            return full + line[len(short):]
    return line


def match_exercise(line, names):
    """Longest exercise name that the line starts with, as whole words."""
    best = None
    for name in names:
        if line == name or (line.startswith(name) and not line[len(name)].isalnum()):
            if best is None or len(name) > len(best):
                best = name
    return best


def split_note(text):
    """Pull '(...)' fragments out as a note."""
    notes = re.findall(r"\(([^)]*)\)", text)
    rest = clean(re.sub(r"\([^)]*\)", " ", text))
    return rest, "; ".join(clean(n) for n in notes if clean(n))


class Session:
    """What was done in one sitting (own practice or a lesson): one entry per exercise."""

    def __init__(self):
        self.entries = OrderedDict()  # exercise -> {"tags": [...], "notes": [...]}
        self.notes = []
        self.last = None              # (exercise, index of the tag a second line extends)
        self.seen = defaultdict(int)  # spreadsheet names, to spot one written twice

    def add(self, exercise, tags=(), note=""):
        e = self.entries.setdefault(exercise, {"tags": [], "notes": []})
        for t in tags:
            t = clean(t)
            if t and t not in e["tags"]:
                e["tags"].append(t)
        if note:
            e["notes"].append(note)
        return e

    def extend(self, text):
        """A second line ("then 37s", "down"): part of the variation above."""
        if not self.last:
            return False
        exercise, idx = self.last
        e = self.entries[exercise]
        if idx is not None and idx < len(e["tags"]):
            e["tags"][idx] = clean(f"{e['tags'][idx]} {text}")
        else:
            e["tags"].append(text)
            self.last = (exercise, len(e["tags"]) - 1)
        return True


def breathing_line(line, s):
    """Read a breathing line into the session. Returns True if it was one."""
    patterns = [
        (r"all breathing", ["set order", "random order"]),
        (r"sfch no random", ["set order"]),
        (r"sfch songs", ["songs"]),
        (r"sfch 5 sfch", ["set order", "5 sfch"]),
        (r"5 sfch", ["5 sfch"]),
        (r"sfch", ["set order"]),
    ]
    for pat, tags in patterns:
        m = re.match(rf"^{pat}\b(.*)$", line)
        if not m:
            continue
        s.add(SFCH, tags)
        s.last = (SFCH, None)
        rest = m.group(1).strip(" ,")
        if rest and SECONDS.match(rest):
            s.add(HOLD, [rest])
            s.last = (HOLD, len(s.entries[HOLD]["tags"]) - 1)
        elif rest:
            s.add(SFCH, note=rest)
        return True
    if SECONDS.match(line):
        if HOLD in s.entries and s.entries[HOLD]["tags"] and line.lower().startswith("then"):
            s.last = (HOLD, len(s.entries[HOLD]["tags"]) - 1)
            return s.extend(line)
        s.add(HOLD, [re.sub(r"^then\s+", "", line)])
        s.last = (HOLD, len(s.entries[HOLD]["tags"]) - 1)
        return True
    return False


def read_cell(raw, names, s, report, date):
    """Read one cell into session s; return notes that weren't exercises."""
    key = str(raw).strip()
    if key in CELL_OVERRIDES:
        for ex, tags, note in CELL_OVERRIDES[key]:
            s.add(ex, tags, note)
        report.add("Hand-made readings (CELL_OVERRIDES)", date, key,
                   " + ".join(ex for ex, _, _ in CELL_OVERRIDES[key]) or "ignored")
        return []
    notes = []
    lines = [clean(l) for l in str(raw).split("\n") if clean(l)]
    for i, line in enumerate(lines):
        fixed = apply_aliases(line)
        if fixed != line:
            report.add("Typos and shorthand corrected", date, line, fixed)
        body, note = split_note(fixed)
        if not body and note and s.last:  # a line that is only "(...)": a note on the line above
            s.add(s.last[0], note=note)
            continue
        if breathing_line(body, s):
            if note:
                s.add(s.last[0], note=note)
            continue
        name = match_exercise(body, names)
        if name:
            s.seen[name] += 1
            own = body[len(name):].strip(" ,-")
            target, tag = COMBINE.get(name, (name, None))
            if tag:
                report.add("Combined into another exercise (COMBINE)", date, name, f"{target} [{tag}]")
            tags = ([tag] if tag else []) + ([own] if own else [])
            if not tags and target in PLAIN_MEANS:
                tags = [PLAIN_MEANS[target]]
            e = s.add(target, tags, note)
            s.last = (target, e["tags"].index(own) if own else None)
        elif i > 0 and body and s.extend(body):
            ex, idx = s.last
            report.add("Second lines joined to the variation above", date, str(raw).replace("\n", " / "),
                       f"{ex} {s.entries[ex]['tags']}")
            if note:
                s.add(ex, note=note)
        else:
            notes.append(clean(line))
    return notes


def date_of(ws, row, prev_date):
    value = ws.cell(row, DAY_COL).value
    if isinstance(value, dt.datetime):
        return value.date()
    if isinstance(value, dt.date):
        return value
    if prev_date is None:
        raise ValueError(f"Row {row}: can't read date {value!r}")
    return prev_date + dt.timedelta(days=1)


def split_row(cells, lesson_from=None):
    """Own practice = from B up to the first empty cell. The rest was at a lesson."""
    filled = [i for i, v in enumerate(cells) if v is not None and clean(v)]
    if not filled:
        return [], []
    if lesson_from is not None:
        return ([cells[i] for i in filled if i < lesson_from], [cells[i] for i in filled if i >= lesson_from])
    if filled[0] >= 3:  # starts at E or later: a lesson-only day
        return [], [cells[i] for i in filled]
    own, lesson, gap = [], [], False
    for i in range(filled[0], len(cells)):
        v = cells[i]
        if v is None or not clean(v):
            gap = True
            continue
        (lesson if gap else own).append(v)
    return own, lesson


def build(path, today=None, stamp=None):
    import openpyxl

    today = today or dt.date.today()
    wb = openpyxl.load_workbook(path, data_only=True)
    ws = wb.active
    report = Report()
    stamp = stamp or dt.datetime.now(dt.timezone.utc).strftime("%Y-%m-%dT%H:%M:%S.000Z")

    # Pass 1: dates, and the rotation pool (for the exercise list and its order).
    rows, prev = [], None
    for r in range(1, ws.max_row + 1):
        if ws.cell(r, DAY_COL).value is None and not any(ws.cell(r, c).value for c in LOG_COLS):
            continue
        d = date_of(ws, r, prev)
        if not isinstance(ws.cell(r, DAY_COL).value, (dt.date, dt.datetime)):
            report.add("Dates that weren't real dates", str(d), str(ws.cell(r, DAY_COL).value),
                       f"read as {d} (the day after the row above)")
        prev = d
        rows.append((r, d))

    last_logged = max(d for r, d in rows if any(ws.cell(r, c).value is not None for c in LOG_COLS))
    recent_from = last_logged - dt.timedelta(weeks=RECENT_WEEKS)
    pool_slot, pool_recent, pool_all = {}, set(), set()
    for r, d in rows:
        for c in POOL_COLS:
            v = ws.cell(r, c).value
            if v is None or not clean(v):
                continue
            name = clean(v)
            pool_all.add(name)
            pool_slot[name] = ((d.weekday() + 1) % 7, c)  # Sunday first, like the sheet
            if d >= recent_from:
                pool_recent.add(name)

    names = sorted(pool_all | set(EXTRA_EXERCISES) | set(COMBINE) | {t for t, _ in COMBINE.values()},
                   key=len, reverse=True)

    # Pass 2: the log.
    entries, days, used = [], [], defaultdict(int)
    for r, d in rows:
        if d > today:
            continue
        cells = [ws.cell(r, c).value for c in LOG_COLS]
        own_cells, lesson_cells = split_row(cells, LESSON_STARTS.get(d.isoformat()))
        if not own_cells and not lesson_cells:
            continue
        date = d.isoformat()
        # Each cell is one go at an exercise, so the same exercise twice in a day (or once
        # on your own and again at the lesson) stays as separate entries.
        day_notes, label, done, seen_own = [], "", [], None
        for lesson, cells_ in [(False, own_cells), (True, lesson_cells)]:
            seen = defaultdict(int)
            got = []
            for v in cells_:
                s = Session()
                s.seen = seen
                day_notes += read_cell(str(v), names, s, report, date)
                got += [(lesson, ex, e) for ex, e in s.entries.items()]
            done += got
            if lesson and cells_:
                report.add("Lesson cells (after a gap in the row)", date,
                           " | ".join(clean(v) for v in cells_), ", ".join(ex for _, ex, _ in got) or "notes only")
            if not lesson:
                seen_own = seen
        if len(own_cells) + len(lesson_cells) == 1 and not done:
            label, day_notes = clean((own_cells + lesson_cells)[0]), []
        for n in day_notes:
            report.add("Cells read as a note on the day", date, n, "day note")

        twice = [n for n, k in seen_own.items() if k > 1]
        if twice:
            report.add("Same exercise twice with no gap: maybe the second was at a lesson (read as own practice)",
                       date, " | ".join(clean(v) for v in own_cells), ", ".join(twice))
        n = 0
        for lesson, ex, e in done:
            n += 1
            used[ex] += 1
            rec = {"id": f"h-{date}-{n:02d}", "date": date, "exercise": slug(ex), "variations": e["tags"]}
            if lesson:
                rec["lesson"] = True
            rec["note"] = "; ".join(e["notes"])
            rec["updated"] = stamp
            entries.append(rec)
        if label or day_notes:
            days.append({"id": date, "label": label, "note": "; ".join(day_notes), "updated": stamp})

    # Exercises: breathing, then the rotation in its grid order, then archived ones.
    def target(name):
        return COMBINE.get(name, (name, None))[0]

    members = defaultdict(set)
    for name in pool_all | set(EXTRA_EXERCISES):
        members[target(name)].add(name)
    for t, _ in COMBINE.values():
        members[t].add(t)
    active = {t for t, m in members.items() if m & pool_recent}
    slot = {t: min((pool_slot[m] for m in ms if m in pool_slot and m in pool_recent), default=(9, 99))
            for t, ms in members.items()}
    ordered_names = BREATHING + sorted(active - set(BREATHING), key=lambda t: slot[t]) + \
        sorted(set(members) - active - set(BREATHING))

    exercises, ids = [], {}
    for order, name in enumerate(ordered_names, 1):
        ex_id = slug(name)
        if ex_id in ids.values():
            raise ValueError(f"Two exercises share the id {ex_id!r}; add one to ID_OVERRIDES")
        ids[name] = ex_id
        archived = name not in BREATHING and name not in active
        exercises.append({
            "id": ex_id, "name": name,
            "category": "breathing" if name in BREATHING else "exercises",
            "variations": VARIATIONS.get(name, []),
            "archived": archived, "order": order, "updated": stamp,
        })
        if archived:
            report.add("Archived (not in the pool in the last %d weeks)" % RECENT_WEEKS, "", name,
                       f"{used.get(name, 0)} entries")
    unknown = {e["exercise"] for e in entries} - set(ids.values())
    if unknown:
        raise ValueError(f"Entries point at exercises that weren't created: {sorted(unknown)}")

    categories = [{"id": cid, "name": cname, "order": i, "updated": stamp}
                  for i, (cid, cname) in enumerate(CATEGORIES, 1)]
    doc = {"app": "practice-log", "version": 2, "categories": categories,
           "exercises": exercises, "days": days, "entries": entries}
    return doc, report


def ordered(rec):
    rank = {k: i for i, k in enumerate(KEY_ORDER)}
    return {k: rec[k] for k in sorted(rec, key=lambda k: (rank.get(k, 99), k))}


def serialize(doc):
    """The same text the app writes: one record per line, so git diffs stay readable."""
    def block(key, items):
        if not items:
            return f'  "{key}": []'
        lines = ",\n".join("    " + json.dumps(ordered(x), ensure_ascii=False, separators=(",", ":")) for x in items)
        return f'  "{key}": [\n{lines}\n  ]'
    parts = [
        block("categories", sorted(doc["categories"], key=lambda x: x["id"])),
        block("exercises", sorted(doc["exercises"], key=lambda x: x["id"])),
        block("days", sorted(doc["days"], key=lambda x: x["id"])),
        block("entries", sorted(doc["entries"], key=lambda x: (x["date"], x["id"]))),
    ]
    return '{\n  "app": "practice-log",\n  "version": 2,\n' + ",\n".join(parts) + "\n}\n"


def main(argv=None):
    ap = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    ap.add_argument("xlsx")
    ap.add_argument("-o", "--out", default="log.json")
    ap.add_argument("--today", help="ignore rows after this date (YYYY-MM-DD); default today")
    ap.add_argument("--stamp", help="the 'updated' time to give every record (default: now). Set it just after "
                    "the time in a file you already loaded, and loading the new one replaces the old import "
                    "without undoing anything you changed in the app since.")
    args = ap.parse_args(argv)
    today = dt.date.fromisoformat(args.today) if args.today else None
    doc, report = build(args.xlsx, today, args.stamp)
    with open(args.out, "w", encoding="utf-8", newline="\n") as f:
        f.write(serialize(doc))
    practice_days = len({e["date"] for e in doc["entries"]})
    lessons = len({e["date"] for e in doc["entries"] if e.get("lesson")})
    active = sum(1 for e in doc["exercises"] if not e["archived"])
    print(f"Wrote {args.out}: {len(doc['entries'])} entries over {practice_days} days "
          f"({lessons} days with a lesson), {len(doc['days'])} days with a label or note, "
          f"{len(doc['exercises'])} exercises ({active} on the board).")
    print(report.render())


if __name__ == "__main__":
    sys.exit(main())
